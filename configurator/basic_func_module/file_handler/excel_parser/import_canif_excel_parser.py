# This Python file uses the following encoding: utf-8
############################################################################
# Copyright (c) 2025 Li Auto Inc. and its affiliates
# Licensed under the Apache License, Version 2.0(the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
############################################################################
import os
from basic_func_module.file_handler.excel_parser.import_excel_base import ImportExcelBase
from basic_func_module.base_logger import logger
from openpyxl.reader.excel import load_workbook


class ImportCanIfExcelParser(ImportExcelBase):

    def execute_import_excel_command(self, excel_path_dict):
        dict = {}
        if not excel_path_dict:
            return None
        for excel_path in excel_path_dict:
            result = self.parse_excel_file(excel_path)
            if result is None:
                logger.error(f"{excel_path} has invalid info!")
                return None

            if self.excel_validate(result) is False:
                logger.info(f"{excel_path} has invalid info!")
                return None
            excel_name = os.path.basename(excel_path)
            dict.update({excel_name: result})

        return dict

    def parse_excel_file(self, file_path) -> dict:
        canif_excel_dict = {}
        workbook = load_workbook(file_path)
        sheet = workbook["Sheet1"]
        for row in range(2, sheet.max_row + 1):
            row_value_list = self.get_row_value(sheet, row)
            canif_excel_dict.update({row_value_list[0]: row_value_list[1]})

        return canif_excel_dict

    def excel_validate(self, dict: dict) -> bool:
        validate_result = True

        for key, value in dict.items():
            pass

        return validate_result
