# This Python file uses the following encoding: utf-8
############################################################################
# Copyright (c) 2025 Li Auto Inc. and its affiliates
# Licensed under the Apache License, Version 2.0(the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
############################################################################
import os
import sys
import copy
import subprocess
from openpyxl import Workbook
from basic_func_module.autosar_utils.autosar_util import AutosarUtil
from func_module.health_monitor.logger import logger
from func_module.mapping.schedule_table_mapping import ScheduleTableManager
from func_module.mapping.task_mapping_info import TaskMappingInfo
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from func_module.health_monitor import perf_record

CUR_FILE_PATH = os.path.dirname(__file__)
TASK_MAPPING = "task_mapping"


class ExportFile:

    def __init__(self):
        pass

    @classmethod
    def parse_os_info(self, data_management):
        os_autosar = data_management.autosar_bsw_class.get("Os")
        rte_autosar = data_management.autosar_bsw_class.get("Rte")
        asw_autosar = data_management.autosar_asw_class
        ib_bswmds = data_management.autosar_ib_class

        # 反推调度表
        sch_info_dict = ScheduleTableManager.re_reasoning_all(os_autosar, rte_autosar, asw_autosar, ib_bswmds)
        # runnable
        ori_runnable_info_dict = TaskMappingInfo.get_runnable_by_trigger_event(rte_autosar, asw_autosar, ib_bswmds)
        runnable_info_dict = {"non_timing": ori_runnable_info_dict["non_timing"], "timing": {"sch": {}, "non_sch": {}}}
        # 调度表启动的周期runnable
        from basic_func_module.base_common.load_plug_in import LoadPlugIn

        os_process = LoadPlugIn.load_module("Rte", "config_plug_in.bsw_process.os_process")
        OsManager = getattr(os_process, "OsManager")
        for sch in sch_info_dict.values():
            task_offset_table = {}
            for offset, task_list in sch["task_offset_dict"].items():
                for task in task_list:
                    task_offset_table[task] = offset
            counter_info = OsManager.get_counter_info(os_autosar, sch["counter_ref"])
            task_offset_table = TaskMappingInfo.convert_tick_to_microseconds(task_offset_table, counter_info)

            sch_runnable_info = copy.deepcopy(ori_runnable_info_dict["timing"])
            sch_runnable_info = {k: sch_runnable_info[k] for k in task_offset_table.keys() if k in sch_runnable_info}
            for task_name in sch_runnable_info.keys():
                for r in sch_runnable_info[task_name]:
                    r["sch"] = sch["name"]
                    r["task_offset"] = task_offset_table[task_name]
            sch_runnable_in_ms = TaskMappingInfo.add_task_offset_to_runnable(sch_runnable_info, task_offset_table)

            for task_name in sch_runnable_in_ms.keys():
                runnable_info_dict["timing"]["sch"][task_name] = sch_runnable_in_ms[task_name]
        # 非调度表启动的周期runnable
        non_sch_runnable_info = {
            key: value
            for key, value in ori_runnable_info_dict["timing"].items()
            if key not in runnable_info_dict["timing"]["sch"].keys()
        }
        runnable_info_dict["timing"]["non_sch"] = non_sch_runnable_info

        # application/task/ISR
        application_info_dict, isr_info_dict, task_info_dict = AutosarUtil.parse_os_info(os_autosar)
        task_mapping_info: dict = {}
        for app, info in application_info_dict.items():
            core_name = info["core_name"]
            task_mapping_info.setdefault(core_name, {})
            task_mapping_info[core_name][app] = {"isr": {}, "task": {}}
            for isr_name in info["isr_name_list"]:
                task_mapping_info[core_name][app]["isr"][isr_name] = isr_info_dict[isr_name]
            sorted_keys = sorted(
                task_mapping_info[core_name][app]["isr"],
                key=lambda k: task_mapping_info[core_name][app]["isr"][k]["OsIsrInterruptPriority"],
                reverse=True,
            )
            task_mapping_info[core_name][app]["isr"] = {
                k: task_mapping_info[core_name][app]["isr"][k] for k in sorted_keys
            }
            for task_name in info["task_name_list"]:
                task_mapping_info[core_name][app]["task"][task_name] = task_info_dict[task_name]
            sorted_keys = sorted(
                task_mapping_info[core_name][app]["task"],
                key=lambda k: task_mapping_info[core_name][app]["task"][k]["OsTaskPriority"],
                reverse=True,
            )
            task_mapping_info[core_name][app]["task"] = {
                k: task_mapping_info[core_name][app]["task"][k] for k in sorted_keys
            }
            for task_name in task_mapping_info[core_name][app]["task"]:
                task_mapping_info[core_name][app]["task"][task_name]["runnable"] = {
                    "non_timing": [],
                    "timing": {"sch": [], "non_sch": []},
                }
                if task_name in runnable_info_dict["non_timing"].keys():
                    task_mapping_info[core_name][app]["task"][task_name]["runnable"]["non_timing"] = runnable_info_dict[
                        "non_timing"
                    ][task_name]
                if task_name in runnable_info_dict["timing"]["sch"].keys():
                    task_mapping_info[core_name][app]["task"][task_name]["runnable"]["timing"]["sch"] = (
                        runnable_info_dict["timing"]["sch"][task_name]
                    )
                elif task_name in runnable_info_dict["timing"]["non_sch"].keys():
                    task_mapping_info[core_name][app]["task"][task_name]["runnable"]["timing"]["non_sch"] = (
                        runnable_info_dict["timing"]["non_sch"][task_name]
                    )

        return task_mapping_info

    @classmethod
    @perf_record.log_function_stats(["RTE", "OS全景图导出"])
    def export_os_info_to_excel(cls, file_path, is_start_file, data_management):
        task_mapping_info = cls.parse_os_info(data_management)

        workbook = Workbook()
        sheet = workbook.active
        # 表头行列号[开始行, 结束行, 开始列, 结束列]
        header_info = {
            "Core": [1, 2, 1, 1],
            "Application": [1, 2, 2, 2],
            "ISR": [1, 1, 3, 4],
            "ISR_name": [2, 2, 3, 3],
            "OsIsrInterruptPriority": [2, 2, 4, 4],
            "Task": [1, 1, 5, 9],
            "Task_name": [2, 2, 5, 5],
            "OsTaskType": [2, 2, 6, 6],
            "OsTaskSchedule": [2, 2, 7, 7],
            "OsTaskPriority": [2, 2, 8, 8],
            "OsTaskStackSize(Byte)": [2, 2, 9, 9],
            "Runnable": [1, 1, 11, 11],
            "Runnable_name": [1, 2, 10, 10],
            "position": [2, 2, 11, 11],
            "activator": [2, 2, 12, 12],
            "period(us)": [2, 2, 13, 13],
            "offset(us)": [2, 2, 14, 14],
            "task_offset_by_schtbl(us)": [2, 2, 15, 15],
        }
        cls.write_merged_header(sheet, header_info)

        cur_row = 3
        for core_name, core_info in task_mapping_info.items():
            core_row = cur_row
            for app_name, app_info in core_info.items():
                app_row = cur_row
                for isr_name, isr_info in app_info["isr"].items():
                    cls.write_sheet_cell(sheet, cur_row, header_info["ISR_name"][2], isr_name)
                    cls.write_sheet_cell(
                        sheet, cur_row, header_info["OsIsrInterruptPriority"][2], isr_info["OsIsrInterruptPriority"]
                    )
                    cur_row += 1
                for task_name, task_info in app_info["task"].items():
                    task_row = cur_row
                    runnable_info = task_info["runnable"]
                    for r in runnable_info["timing"]["sch"]:
                        cls.write_sheet_cell(sheet, cur_row, header_info["Runnable_name"][2], r["runnable_name"])
                        cls.write_sheet_cell(sheet, cur_row, header_info["position"][2], r["position"])
                        cls.write_sheet_cell(sheet, cur_row, header_info["activator"][2], f"(schtbl){r['sch']}")
                        cls.write_sheet_cell(sheet, cur_row, header_info["period(us)"][2], r["runnable_period"])
                        cls.write_sheet_cell(sheet, cur_row, header_info["offset(us)"][2], r["runnable_offset"])
                        cls.write_sheet_cell(
                            sheet, cur_row, header_info["task_offset_by_schtbl(us)"][2], r["task_offset"]
                        )
                        cur_row += 1
                    for r in runnable_info["timing"]["non_sch"]:
                        cls.write_sheet_cell(sheet, cur_row, header_info["Runnable_name"][2], r["runnable_name"])
                        cls.write_sheet_cell(sheet, cur_row, header_info["position"][2], r["position"])
                        cls.write_sheet_cell(
                            sheet, cur_row, header_info["activator"][2], f"(alarm){r['alarm_ref'].split('/')[-1]}"
                        )
                        cls.write_sheet_cell(sheet, cur_row, header_info["period(us)"][2], r["runnable_period"])
                        cls.write_sheet_cell(sheet, cur_row, header_info["offset(us)"][2], r["runnable_offset"])
                        cls.write_sheet_cell(sheet, cur_row, header_info["task_offset_by_schtbl(us)"][2], "N/A")
                        cur_row += 1
                    for r in runnable_info["non_timing"]:
                        cls.write_sheet_cell(sheet, cur_row, header_info["Runnable_name"][2], r["runnable_name"])
                        cls.write_sheet_cell(sheet, cur_row, header_info["position"][2], r["position"])
                        cur_row += 1
                    cls.write_sheet_cell(sheet, task_row, header_info["Task_name"][2], task_name)
                    cls.write_sheet_cell(sheet, task_row, header_info["OsTaskType"][2], task_info["OsTaskType"])
                    cls.write_sheet_cell(sheet, task_row, header_info["OsTaskSchedule"][2], task_info["OsTaskSchedule"])
                    cls.write_sheet_cell(sheet, task_row, header_info["OsTaskPriority"][2], task_info["OsTaskPriority"])
                    cls.write_sheet_cell(
                        sheet, task_row, header_info["OsTaskStackSize(Byte)"][2], task_info["OsTaskStackSize"]
                    )
                    if task_row == cur_row:
                        cur_row += 1
                    sheet.merge_cells(
                        start_row=task_row,
                        end_row=cur_row - 1,
                        start_column=header_info["Task_name"][2],
                        end_column=header_info["Task_name"][2],
                    )
                    sheet.merge_cells(
                        start_row=task_row,
                        end_row=cur_row - 1,
                        start_column=header_info["OsTaskType"][2],
                        end_column=header_info["OsTaskType"][2],
                    )
                    sheet.merge_cells(
                        start_row=task_row,
                        end_row=cur_row - 1,
                        start_column=header_info["OsTaskSchedule"][2],
                        end_column=header_info["OsTaskSchedule"][2],
                    )
                    sheet.merge_cells(
                        start_row=task_row,
                        end_row=cur_row - 1,
                        start_column=header_info["OsTaskPriority"][2],
                        end_column=header_info["OsTaskPriority"][2],
                    )
                    sheet.merge_cells(
                        start_row=task_row,
                        end_row=cur_row - 1,
                        start_column=header_info["OsTaskStackSize(Byte)"][2],
                        end_column=header_info["OsTaskStackSize(Byte)"][2],
                    )
                cls.write_sheet_cell(sheet, app_row, header_info["Application"][2], app_name)
                if app_row == cur_row:
                    cur_row += 1
                sheet.merge_cells(
                    start_row=app_row,
                    end_row=cur_row - 1,
                    start_column=header_info["Application"][2],
                    end_column=header_info["Application"][2],
                )
            cls.write_sheet_cell(sheet, core_row, header_info["Core"][2], core_name)
            if core_row == cur_row:
                cur_row += 1
            sheet.merge_cells(
                start_row=core_row,
                end_row=cur_row - 1,
                start_column=header_info["Core"][2],
                end_column=header_info["Core"][2],
            )

        cls.set_column_widths(sheet, ori_length=5)
        cls.set_row_heights(sheet)

        color_dict = {
            0: PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
            1: PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
            2: PatternFill(start_color="ccffff", end_color="ccffff", fill_type="solid"),
            3: PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
            4: PatternFill(start_color="C6E2FF", end_color="C6E2FF", fill_type="solid"),
            5: PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
            6: PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
            7: PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
            8: PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
            9: PatternFill(start_color="FFE7BA", end_color="FFE7BA", fill_type="solid"),
            10: PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
            11: PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
            12: PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
            13: PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
            14: PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
        }
        cls.set_column_color(sheet, color_dict)

        try:
            workbook.save(file_path)
            if is_start_file is True:
                if sys.platform == "win32":
                    os.startfile(file_path)
                elif sys.platform == "linux":
                    subprocess.call(f"nohup libreoffice {file_path} &", shell=True)
        except PermissionError:
            logger.error("File is open in another application. Please close it and try again.")

    @classmethod
    def write_merged_header(cls, sheet, header_info):
        font = Font(name="Cascadia Mono", size=12, color="FFFFFF", bold=True, italic=False, strike=False)
        border = Border(
            left=Side(border_style="thin", color="000000"),
            right=Side(border_style="thin", color="000000"),
            top=Side(border_style="thin", color="000000"),
            bottom=Side(border_style="thin", color="000000"),
        )
        alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for name, idx_list in header_info.items():
            cell = sheet.cell(row=idx_list[0], column=idx_list[2], value=name)
            cell.font = font
            cell.fill = PatternFill("solid", fgColor="3A5FCD")
            cell.border = border
            cell.alignment = alignment
            if idx_list[0] != idx_list[1] or idx_list[2] != idx_list[3]:
                sheet.merge_cells(
                    start_row=idx_list[0], end_row=idx_list[1], start_column=idx_list[2], end_column=idx_list[3]
                )

    @classmethod
    def write_sheet_cell(cls, sheet, row, column, value):
        cell = sheet.cell(row=row, column=column, value=value)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    @classmethod
    def mapping_info_export_to_excel(cls, info_table, mapping_type, file_path, data_management, is_start_file=True):
        if "export_os_info" == mapping_type:
            cls.export_os_info_to_excel(file_path, is_start_file, data_management)
            return

        workbook = Workbook()
        sheet = workbook.active

        table = info_table

        if TASK_MAPPING == mapping_type:
            header = cls.get_headers_model(table.table_headers)
            item_list = table.item_list
            column_count = len(header)
            if column_count == 0:
                row_count = 0
            else:
                row_count = int(len(item_list) / column_count)
        else:
            header = cls.get_headers(table)
            row_count = table.rowCount()
            column_count = table.columnCount()

        cls.write_header(sheet, header)
        if TASK_MAPPING == mapping_type:
            cls.write_data_model(column_count, item_list, sheet)
        else:
            cls.write_data(row_count, column_count, table, sheet)
        cls.set_column_widths(sheet)
        cls.set_row_heights(sheet)

        try:
            workbook.save(file_path)
            if is_start_file is True:
                if sys.platform == "win32":
                    os.startfile(file_path)
                elif sys.platform == "linux":
                    subprocess.call(f"nohup libreoffice {file_path} &", shell=True)
        except PermissionError:
            logger.error("File is open in another application. Please close it and try again.")

    @classmethod
    def get_headers(cls, table):
        header_data = []
        for column in range(table.columnCount()):
            item = table.headerData(column)
            if item is not None:
                header_data.append(item)
            else:
                header_data.append("")

        return header_data

    @classmethod
    def get_headers_model(cls, table_headers):
        header_data = []
        for i in range(len(table_headers)):
            header_data.append(
                f"{table_headers[i]['OsApplName']}\n"
                f"{table_headers[i]['OsTaskName']}\n"
                f"Priority {table_headers[i]['OsTaskPriority']}"
            )

        return header_data

    @classmethod
    def write_header(cls, sheet, header):
        # Set header style
        header_font = Font(bold=True)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header_row = []

        for value in header:
            cell = sheet.cell(row=1, column=len(header_row) + 1, value=value)
            cell.font = header_font
            cell.alignment = header_alignment
            header_row.append(cell)

    @classmethod
    def write_data(cls, row_count, column_count, table, sheet):
        data_font = Font()
        data_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        for row in range(row_count):
            for column in range(column_count):
                item = table.index(row, column)
                if item is not None:
                    value = item.data()
                    cell = sheet.cell(row=row + 2, column=column + 1, value=value)
                    cell.font = data_font
                    cell.alignment = data_alignment

    @classmethod
    def write_data_model(cls, column_count, item_lists, sheet):
        data_font = Font()
        data_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        row_index = 2
        col_index = 1
        for item in item_lists:
            if item.data is None:
                data = ""
            else:
                data = str(item.data["RunnableName"])
            cell = sheet.cell(row=row_index, column=col_index, value=data)
            col_index += 1
            if col_index > column_count:
                row_index += 1
                col_index = 1
            cell.font = data_font
            cell.alignment = data_alignment

    @classmethod
    def set_row_heights(cls, sheet):
        for row in sheet.rows:
            row_height = 18
            for cell in row:
                try:
                    lines = str(cell.value).count("\n") + 1
                    adjusted_height = lines * 15
                    if adjusted_height > row_height:
                        row_height = adjusted_height
                except BaseException:
                    pass
            sheet.row_dimensions[row[0].row].height = row_height

    @classmethod
    def set_column_widths(cls, sheet, ori_length=0):

        for column in sheet.columns:
            max_length = ori_length
            for _ in column:
                if not hasattr(_, "column_letter"):
                    continue
                column_name = _.column_letter
                break
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except BaseException:
                    pass
            adjusted_width = (max_length + 2) * 1.2
            sheet.column_dimensions[column_name].width = adjusted_width

    @classmethod
    def set_column_color(cls, sheet, color_dict):
        border = Border(
            left=Side(border_style="thin", color="000000"),
            right=Side(border_style="thin", color="000000"),
            top=Side(border_style="thin", color="000000"),
            bottom=Side(border_style="thin", color="000000"),
        )
        i = -1
        for column in sheet.columns:
            i += 1
            if i not in color_dict:
                continue
            color = color_dict[i]
            for cell in column[2:]:
                cell.fill = color
                cell.border = border
