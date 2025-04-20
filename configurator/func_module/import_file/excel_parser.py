# This Python file uses the following encoding: utf-8
############################################################################
# Copyright (c) 2025 Li Auto Inc. and its affiliates
# Licensed under the Apache License, Version 2.0(the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
############################################################################
import openpyxl


class TaskMappingItem(object):

    def __init__(self, runnable_name="", task_name="", position=-1):
        self.runnable_name = runnable_name
        self.task_name = task_name
        self.position = position


class ExcelParser:
    excel_used_sheet_name = "Sheet"
    excel_title_row_index = 1

    taskmapping_item_list = []
    log_list = []

    @classmethod
    def excel_2_taskmapping_item_list(cls, excel_file):
        cls.taskmapping_item_list = []
        cls.log_list = []

        try:
            wb = openpyxl.load_workbook(excel_file, data_only=True)
        except Exception as e:
            raise Exception("parse Interface Excel Data Failed due to \n{0}".format(e))

        sh = wb[cls.excel_used_sheet_name]
        for column in range(sh.min_column, sh.max_column + 1):
            table_header = sh.cell(cls.excel_title_row_index, column).value
            if not table_header:
                continue
            table_content_list = table_header.split("\n")
            if len(table_content_list) != 3:
                log_str = f"excel file column:{column} format error!"
                print(log_str)
                cls.log_list.append(log_str)
                continue

            task_name = table_content_list[1]
            for row in range(cls.excel_title_row_index + 1, sh.max_row + 1):
                if not sh.cell(row, column).value:
                    continue
                runnable_name = sh.cell(row, column).value.strip()
                task_mapping_item = TaskMappingItem(runnable_name, task_name, row - 2)
                cls.taskmapping_item_list.append(task_mapping_item)

        return cls.taskmapping_item_list

    @classmethod
    def get_taskmapping_item(cls, runnable_name) -> TaskMappingItem:
        for taskmapping_item in cls.taskmapping_item_list:
            if taskmapping_item.runnable_name == runnable_name:
                return taskmapping_item
        return None
