# This Python file uses the following encoding: utf-8
############################################################################
# Copyright (c) 2025 Li Auto Inc. and its affiliates
# Licensed under the Apache License, Version 2.0(the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
############################################################################
import os
from enum import Enum, IntEnum
from copy import deepcopy
from openpyxl import load_workbook

from func_module.health_monitor.logger import logger
from basic_func_module.autosar_utils.autosar_util import AutosarUtil
from basic_func_module.autosar_utils.autosar_utils import AutosarUtils
from basic_func_module.autosar_utils.utils import (
    get_short_name,
    get_definition_ref,
    get_value,
    get_value_ref,
    get_valueOf,
)
from basic_func_module.autosar_utils.autosar_flatextract import (
    create_sender_receiver_to_signal_mapping,
    create_sender_receiver_to_signal_group_mapping,
)
from func_module.module_management.arxml_handler.system.system_arxmls_manager import SystemArxmlsManager

CUR_FILE_PATH = os.path.dirname(__file__)

FLAT_EXTRACT_PREFIX = "/COM/VCOS2/CFG/WORKFLOW/SYSDESC/SYNC"

_element_key = "Communication Element"
_component_key = "Component Prototype"
_port_key = "Port Prototype"
_signal_key = "System Signal"
_pdu_key = "PDU"


class Compatibility(Enum):
    full = (0, "Full")
    partial = (1, "Partial")
    full_datalength_compatible = (2, "Full & Data Type Length Compatible")
    partial_datalength_compatible = (3, "Partial & Data Type Length Compatible")
    datalength_imcompatible = (4, "Data Type Length Imcompatible")
    imcompatible = (5, "")


PORT_SIGNAL_1TO1_AUTOMAPPING_DATATYPE_SAME = [Compatibility.full, Compatibility.partial]

PORT_SIGNAL_1to1_AUTOMAPPING_DATATYPE_COMPATIBLE = [
    Compatibility.full,
    Compatibility.partial,
    Compatibility.full_datalength_compatible,
    Compatibility.partial_datalength_compatible,
]

PORT_SIGNAL_1TON_AUTOMAPPING_DATATYPE_SAME = [Compatibility.full]

PORT_SIGNAL_1toN_AUTOMAPPING_DATATYPE_COMPATIBLE = [Compatibility.full, Compatibility.full_datalength_compatible]


class DataMapping:

    DATA_TYPE_MATCHING = {
        "BOOLEAN": ["boolean"],
        "SINT8": ["sint8"],
        "UINT8": ["uint8"],
        "SINT16": ["sint16"],
        "UINT16": ["uint16"],
        "SINT32": ["sint32"],
        "UINT32": ["uint32"],
        "FLOAT32": ["float32"],
        "SINT64": ["sint64"],
        "UINT64": ["uint64"],
        "FLOAT64": ["float64"],
        "UINT8_N": [],
        "UINT8_DYN": [],
    }

    NATIVE_DATATYPES_SIZE = {
        "boolean": 1,
        "sint8": 1,
        "uint8": 1,
        "sint16": 2,
        "uint16": 2,
        "sint32": 4,
        "uint32": 4,
        "sint64": 8,
        "uint64": 8,
        "float32": 4,
        "float64": 8,
    }

    class INFOS_INDEX(IntEnum):
        SWC_PORTS = 0
        FLATEXTRACT = 1
        SYSTEM_SIGNALS = 2
        INDEX_NUM = 3

    @classmethod
    def get_datamapping_infos(
        cls, autosar_asw_class, autosar_bsw_class, autosar_system_class_dict, filter_conditions=None
    ):
        infos = [{}] * cls.INFOS_INDEX.INDEX_NUM
        all_swc_ports = cls.get_all_swc_ports_dict(autosar_asw_class, filter_conditions)
        flatextract_datamppings = cls.get_port_to_signal_pairs_from_flatextract(
            autosar_system_class_dict.get("FlatExtract", None)
        )
        system_signals_dict = cls.get_system_signal_to_isignal_mappings_info(
            autosar_system_class_dict.get("Communication", None), autosar_bsw_class.get("Com", None)
        )
        cls.update_data_mappings_to_ports(
            all_swc_ports["all_port_data_s"], flatextract_datamppings, system_signals_dict
        )
        infos[cls.INFOS_INDEX.SWC_PORTS] = all_swc_ports
        infos[cls.INFOS_INDEX.FLATEXTRACT] = flatextract_datamppings
        infos[cls.INFOS_INDEX.SYSTEM_SIGNALS] = system_signals_dict
        return infos

    @classmethod
    def update_flatextract_data_mapping_from_files_import(cls, asw_autosar, system_autosar_dict, com_autosar):
        asw_ports = cls.get_all_swc_ports_dict(asw_autosar)
        system_signals_dict = cls.get_system_signal_to_isignal_mappings_info(
            system_autosar_dict.get("Communication", None), com_autosar
        )
        autosar_flatextract = system_autosar_dict.get("FlatExtract", None)

        if autosar_flatextract is None:
            return
        data_mappings_obj = autosar_flatextract.find(f"{FLAT_EXTRACT_PREFIX}/SYSTEM/MAPPINGS")
        if data_mappings_obj is None:
            return
        signal_mapping_list = autosar_flatextract.find_nodes("SENDER-RECEIVER-TO-SIGNAL-MAPPING")
        for signal_mapping in signal_mapping_list:
            if cls.__data_mapping_need_remove(asw_ports, system_signals_dict, signal_mapping, is_group=False):
                AutosarUtil.remove_node_value(
                    data_mappings_obj.DATA_MAPPINGS.SENDER_RECEIVER_TO_SIGNAL_MAPPING, signal_mapping
                )
        signalgroup_mapping_list = autosar_flatextract.find_nodes("SENDER-RECEIVER-TO-SIGNAL-GROUP-MAPPING")
        for group_mapping in signalgroup_mapping_list:
            if cls.__data_mapping_need_remove(asw_ports, system_signals_dict, group_mapping, is_group=True):
                AutosarUtil.remove_node_value(
                    data_mappings_obj.DATA_MAPPINGS.SENDER_RECEIVER_TO_SIGNAL_GROUP_MAPPING, group_mapping
                )

    @classmethod
    def __data_mapping_need_remove(cls, asw_ports, system_signals_dict, data_mapping, is_group=False):
        if is_group:
            mapping_info = SystemArxmlsManager.paser_signal_group_mapping(data_mapping)
        else:
            mapping_info = SystemArxmlsManager.paser_signal_mapping(data_mapping)
        if mapping_info is None:
            return False
        swc_name = mapping_info["data_element_iref"]["context_component_ref"].split("/")[-1]
        port_name = mapping_info["data_element_iref"]["context_port_ref"].split("/")[-1]
        data_name = mapping_info["data_element_iref"]["target_data_prototype_ref"].split("/")[-1]
        swc_p_d = f"{swc_name}/{port_name}/{data_name}"
        port_info = asw_ports.get("all_port_data_s", {}).get(swc_p_d, None)
        if port_info is None:
            return True
        direction = port_info["direction"]
        signal_type = port_info["allowable_signal_type"]
        if is_group:
            group_name = mapping_info["signal_group_ref"].split("/")[-1]
            group_info = system_signals_dict.get(direction, {}).get(signal_type, {}).get(group_name, None)
            if group_info is None:
                return True
            for member_mapping in mapping_info.get("type_mapping", []):
                signal_name = member_mapping["system_signal_ref"].split("/")[-1]
                if signal_name not in group_info["system_signal"].keys():
                    return True
                record_element = member_mapping["implementation_record_element"]
                if not any(
                    record_element == member["name"]
                    for member in port_info.get("DATA-ELEMENT-REF", {}).get("idt_detail", {}).get("members", [])
                ):
                    return True
        else:
            signal_name = mapping_info["system_signal_ref"].split("/")[-1]
            signal_info = system_signals_dict.get(direction, {}).get(signal_type, {}).get(signal_name, None)
            if signal_info is None:
                return True
        return False

    @classmethod
    def add_new_data_mapping_arxml(cls, autosar_system, mapping_info, signal_type="system_signal"):
        if "FlatExtract" not in autosar_system:
            return False
        data_mappings_obj = autosar_system["FlatExtract"].find(f"{FLAT_EXTRACT_PREFIX}/SYSTEM/MAPPINGS")
        if data_mappings_obj is None:
            logger.error("Can't find data mappings container")
            return False
        if signal_type == "system_signal":
            create_sender_receiver_to_signal_mapping(data_mappings_obj.DATA_MAPPINGS, mapping_info)
        elif signal_type == "system_signal_group":
            create_sender_receiver_to_signal_group_mapping(data_mappings_obj.DATA_MAPPINGS, mapping_info)
        # autosar_system["FlatExtract"] = AutosarUtil.update_autosar(autosar_system["FlatExtract"])
        return True

    @classmethod
    def delete_data_mapping_from_arxml(cls, autosar_system, port_info, signal_info):
        if "FlatExtract" not in autosar_system:
            return False
        data_mappings_obj = autosar_system["FlatExtract"].find(f"{FLAT_EXTRACT_PREFIX}/SYSTEM/MAPPINGS")
        if data_mappings_obj is None:
            return False
        if port_info["allowable_signal_type"] == "system_signal":
            for mapping_obj in data_mappings_obj.DATA_MAPPINGS.SENDER_RECEIVER_TO_SIGNAL_MAPPING:
                mapping_dict = SystemArxmlsManager.paser_signal_mapping(mapping_obj)
                if (
                    mapping_dict
                    and signal_info["system_signal"] == mapping_dict["system_signal_ref"].split("/")[-1]
                    and port_info["Name"] == mapping_dict["data_element_iref"]["context_port_ref"].split("/")[-1]
                    and port_info["DATA-ELEMENT-REF"]["name"]
                    == mapping_dict["data_element_iref"]["target_data_prototype_ref"].split("/")[-1]
                ):
                    AutosarUtil.remove_node_value(
                        data_mappings_obj.DATA_MAPPINGS.SENDER_RECEIVER_TO_SIGNAL_MAPPING, mapping_obj
                    )
                    # autosar_system["FlatExtract"] = AutosarUtil.update_autosar(autosar_system["FlatExtract"])
                    return True
        elif port_info["allowable_signal_type"] == "system_signal_group":
            for mapping_obj in data_mappings_obj.DATA_MAPPINGS.SENDER_RECEIVER_TO_SIGNAL_GROUP_MAPPING:
                mapping_dict = SystemArxmlsManager.paser_signal_group_mapping(mapping_obj)
                if (
                    mapping_dict
                    and signal_info["short_name"] == mapping_dict["signal_group_ref"].split("/")[-1]
                    and port_info["Name"] == mapping_dict["data_element_iref"]["context_port_ref"].split("/")[-1]
                    and port_info["DATA-ELEMENT-REF"]["name"]
                    == mapping_dict["data_element_iref"]["target_data_prototype_ref"].split("/")[-1]
                ):
                    AutosarUtil.remove_node_value(
                        data_mappings_obj.DATA_MAPPINGS.SENDER_RECEIVER_TO_SIGNAL_GROUP_MAPPING, mapping_obj
                    )
                    # autosar_system["FlatExtract"] = AutosarUtil.update_autosar(autosar_system["FlatExtract"])
                    return True
        else:
            return False

    @classmethod
    def __get_data_element_idt(cls, data_element_ref):
        data_element = cls.asw_autosar.find(data_element_ref) if data_element_ref is not None else None
        if data_element is None:
            return None
        type_tref = getattr(data_element.TYPE_TREF, "valueOf_", "")
        if data_element.TYPE_TREF.DEST == "IMPLEMENTATION-DATA-TYPE":
            idt_name = str(type_tref).split("/")[-1]
        else:
            idt_name = str(cls.adt_mappings.get(type_tref, "")).split("/")[-1]
        return idt_name

    @classmethod
    def port_to_mapping_filter(cls, filter_conditions, all_swc_ports):

        def filter_result(port_info, direction):
            if (
                port_info["Name"].startswith(prefix)
                and port_info["Name"].endswith(postfix)
                and (port_info["direction"] == direction or direction is None)
            ):
                return True
            else:
                return False

        rv = []
        prefix = filter_conditions.get("prefix", "")
        postfix = filter_conditions.get("postfix", "")
        direction = filter_conditions.get("direction", None)
        rv = [
            {**port, "port_name_to_mapping": port["Name"][len(prefix) : len(port["Name"]) - len(postfix)]}
            for port in all_swc_ports
            if filter_result(port, direction)
        ]

        return rv

    @classmethod
    def system_signals_to_mapping_filter(cls, filter_conditions, system_signals):
        rv = {"system_signal": {}, "system_signal_group": {}}
        direction = filter_conditions.get("direction", None)
        system_signals_filter_direction = system_signals.get(direction, {})
        case_sensitive = filter_conditions.get("case_sensitive", True)
        if case_sensitive is True:
            return system_signals_filter_direction
        rv["system_signal"] = {
            key.lower(): signal for key, signal in system_signals_filter_direction["system_signal"].items()
        }
        rv["system_signal_group"] = {
            key.lower(): signal for key, signal in system_signals_filter_direction["system_signal_group"].items()
        }
        return rv

    @classmethod
    def get_new_data_mappings_dict(cls, swc_ports, system_signals, filter_conditions):
        rv = {
            "1to1_full": [],
            "1to1_partial": [],
            "1to1_full_datalength_compatible": [],
            "1to1_partial_datalength_compatible": [],
            "1to1_datalength_imcompatible": [],
            "1toN": [],
            "no_mapping": [],
        }
        if filter_conditions["match_compatibility"] is True:
            PORT_SIGNAL_1to1_AUTOMAPPING = PORT_SIGNAL_1to1_AUTOMAPPING_DATATYPE_COMPATIBLE
            PORT_SIGNAL_1toN_AUTOMAPPING = PORT_SIGNAL_1toN_AUTOMAPPING_DATATYPE_COMPATIBLE
        else:
            PORT_SIGNAL_1to1_AUTOMAPPING = PORT_SIGNAL_1TO1_AUTOMAPPING_DATATYPE_SAME
            PORT_SIGNAL_1toN_AUTOMAPPING = PORT_SIGNAL_1TON_AUTOMAPPING_DATATYPE_SAME
        for port in swc_ports:
            if len(port["data_mappings"]) > 0:
                continue
            port["signals_recommend"] = cls._port_to_signal_mapping_recommend(port, system_signals, filter_conditions)
            if len(port["signals_recommend"]) == 1:
                if port["signals_recommend"][0]["compatibility"] in PORT_SIGNAL_1to1_AUTOMAPPING:
                    port["signals_recommend"][0]["auto_mapping"] = True
                else:
                    port["signals_recommend"][0]["auto_mapping"] = False
                if port["signals_recommend"][0]["compatibility"] == Compatibility.full:
                    rv["1to1_full"].append(port)
                elif port["signals_recommend"][0]["compatibility"] == Compatibility.partial:
                    rv["1to1_partial"].append(port)
                elif port["signals_recommend"][0]["compatibility"] == Compatibility.full_datalength_compatible:
                    rv["1to1_full_datalength_compatible"].append(port)
                elif port["signals_recommend"][0]["compatibility"] == Compatibility.partial_datalength_compatible:
                    rv["1to1_partial_datalength_compatible"].append(port)
                elif port["signals_recommend"][0]["compatibility"] == Compatibility.datalength_imcompatible:
                    rv["1to1_datalength_imcompatible"].append(port)
            elif len(port["signals_recommend"]) > 1:
                for signal in port["signals_recommend"]:
                    if signal["compatibility"] in PORT_SIGNAL_1toN_AUTOMAPPING:
                        signal["auto_mapping"] = True
                    else:
                        signal["auto_mapping"] = False
                rv["1toN"].append(port)
            else:
                rv["no_mapping"].append(port)
        return rv

    @classmethod
    def check_data_mapping(cls, port_name, de_name, system_signals, conditions):

        def check_data_mapping_rematch():
            if sg_name.startswith(f"{de_name}_{special_mark}_") or sg_name.startswith(f"{port_name}_{special_mark}_"):
                return Compatibility.full
            else:
                return Compatibility.partial

        special_mark = "vcos"
        compatibility = Compatibility.imcompatible
        mapped_signal = []
        if de_name in system_signals.keys():
            compatibility = Compatibility.full
            mapped_signal.append({"signal": deepcopy(system_signals[de_name]), "compatibility": compatibility})
        elif port_name in system_signals.keys():
            compatibility = Compatibility.full
            mapped_signal.append({"signal": deepcopy(system_signals[port_name]), "compatibility": compatibility})
        else:
            if conditions.get("whole_word") is True and conditions.get("allow_multi_signal", False) is False:
                # 如果开启全词匹配,并不允许1对多连接,后面不需要再匹配,直接返回
                return mapped_signal
            for sg_name, signal in system_signals.items():
                if port_name in sg_name or de_name in sg_name or sg_name in port_name or sg_name in de_name:
                    if conditions.get("allow_multi_signal", False):
                        compatibility = check_data_mapping_rematch()
                    else:
                        compatibility = Compatibility.partial
                    mapped_signal.append({"signal": deepcopy(signal), "compatibility": compatibility})

        if conditions.get("whole_word") is True:
            """
            对匹配到的推荐信号进行处理,如果打开了全词匹配,且所有信号的匹配度都不是full全词匹配,则删除推荐信号,防止误连接
            """
            for signal_info in mapped_signal:
                if signal_info["compatibility"] == Compatibility.full:
                    break
            else:
                mapped_signal = []

        return mapped_signal

    @classmethod
    def get_port_signal_compatibility(cls, port_name, de_name, signal_name):
        if port_name == signal_name or de_name == signal_name:
            return Compatibility.full
        elif port_name in signal_name or de_name in signal_name or signal_name in port_name or signal_name in de_name:
            return Compatibility.partial
        return Compatibility.imcompatible

    @classmethod
    def get_sr_to_signal_mapping_dicts(cls, ports):

        def _add_data_mapping_dict_to_list():
            if port["allowable_signal_type"] == "system_signal":
                data_mapping_dict = cls.get_sr_to_signal_mapping_dict(
                    port["Tag"],
                    port["Component"],
                    port["Port_path"],
                    port["DATA-ELEMENT-REF"]["ref"],
                    signal["signal"]["system_signal"],
                )
            elif port["allowable_signal_type"] == "system_signal_group":
                mapping_info = cls._get_signalgroup_mapping_info(port, signal["signal"])
                data_mapping_dict = cls.get_sr_to_signalgroup_mapping_dict(port, mapping_info)

            rv.append(data_mapping_dict)

        rv = []
        for port in ports:
            for signal in port["signals_recommend"]:
                if signal.get("auto_mapping", False) is True:
                    _add_data_mapping_dict_to_list()
        return rv

    @classmethod
    def get_sr_to_signal_mapping_dict(cls, port_type, component, context_port_ref, data_prototype, signal):
        mapping_info = {
            "data_element_iref": {
                "context_component_ref": f"{FLAT_EXTRACT_PREFIX}/COMPOSITIONTYPE/{component}",
                "context_composition_ref": f"{FLAT_EXTRACT_PREFIX}/SYSTEM/COMPOSITIONTYPE",
                "context_port_ref": context_port_ref,
                "port_type": port_type,
                "target_data_prototype_ref": data_prototype,
            },
            "system_signal_ref": f"/Signal/{signal}",
        }
        return mapping_info

    @classmethod
    def get_sr_to_signalgroup_mapping_dict(cls, port_info, mapping_info):
        sr_to_signalgroup_mapping_dict = {
            "data_element_iref": {
                "context_component_ref": f'{FLAT_EXTRACT_PREFIX}/COMPOSITIONTYPE/{port_info["Component"]}',
                "context_composition_ref": f"{FLAT_EXTRACT_PREFIX}/SYSTEM/COMPOSITIONTYPE",
                "context_port_ref": port_info["Port_path"],
                "port_type": port_info["Tag"],
                "target_data_prototype_ref": port_info["DATA-ELEMENT-REF"]["ref"],
            },
            "signal_group_ref": f'/Signal/{mapping_info["signal_group"]["short_name"]}',
            "type_mapping": mapping_info["type_mapping"],
        }
        # for name, recore_element_mapping in mapping_info["type_mapping"]:
        #     sr_to_signalgroup_mapping_dict["type_mapping"].append({
        #         "implementation_record_element_ref": f'/DataTypes/{port_info["DATA-ELEMENT-REF"]["idt"]}/{name}',
        #         "implementation_record_element": name,
        #         "system_signal_ref": f'/Signal/{recore_element_mapping["signal"]["short_name"]}',
        #         "system_signal": recore_element_mapping["signal"]["short_name"]
        #     })
        return sr_to_signalgroup_mapping_dict

    @classmethod
    def _get_signalgroup_mapping_info(cls, port_info, signal_info):
        rv = {"signal_group": signal_info, "type_mapping": []}
        for _, brief in signal_info["type_mappings_brief"].items():
            element = brief["implementation_record_element"]
            temp_dict = {
                "implementation_record_element_ref": f'/DataTypes/{port_info["DATA-ELEMENT-REF"]["idt"]}/{element}',
                "implementation_record_element": element,
                "system_signal_ref": f'/Signal/{brief["system_signal"]}',
                "system_signal": brief["system_signal"],
            }
            rv["type_mapping"].append(temp_dict)
        return rv

    @classmethod
    def add_new_data_mappings_to_arxml(cls, autosar_system, new_mappings_list):
        ret_list = []
        for new_mapping in new_mappings_list:
            signal_type = "system_signal"
            if "signal_group_ref" in new_mapping.keys():
                signal_type = "system_signal_group"
            ret_list.append(cls.add_new_data_mapping_arxml(autosar_system, new_mapping, signal_type))
        return ret_list

    @classmethod
    def get_all_swc_ports_dict(cls, asw_autosar, filter_conditions=None):
        swc_ports_dict = {"components": {}, "all_port_data_s": {}}
        if asw_autosar is None:
            return swc_ports_dict
        idt_dict, idts_dependency, path_to_idt_mapping = AutosarUtils.get_all_idts_and_dependency(asw_autosar)
        cls.idt_dict = cls.update_idts_basetype(idt_dict)
        mapping_sets, cls.adt_mappings = AutosarUtils.get_datatype_mapping_sets(asw_autosar)
        cls.asw_autosar = asw_autosar
        componenttypes = []
        for swc_type in AutosarUtils.SUPPORT_APPL_COMPONENT_TYPE:
            componenttypes += asw_autosar.find_nodes(swc_type)
        for componenttype in componenttypes:
            direction_filter = None
            if filter_conditions:
                components_filter = filter_conditions.get("components", [])
                swc_name = get_short_name(componenttype)
                if swc_name not in components_filter:
                    continue
                else:
                    direction_filter = filter_conditions.get("direction", None)
            swc_ports, port_data_s = cls.get_swc_ports_dict(componenttype, cls.idt_dict, direction_filter)
            if len(port_data_s) > 0:
                swc_ports_dict["components"].update(swc_ports)
                swc_ports_dict["all_port_data_s"].update(port_data_s)
        swc_ports_dict["components"] = dict(sorted(swc_ports_dict["components"].items()))
        swc_ports_dict["all_port_data_s"] = dict(sorted(swc_ports_dict["all_port_data_s"].items()))
        return swc_ports_dict

    @classmethod
    def get_swc_ports_dict(cls, component, idt_dict, direction_filter=None):
        all_port_data_s = {}
        swc_name = get_short_name(component)
        swc_ports = {swc_name: {"pports": [], "rports": []}}
        component_path = AutosarUtil.get_path(component)
        ports_obj = component.PORTS
        if not ports_obj:
            return swc_ports, all_port_data_s

        port_obj_list = []
        if direction_filter == "Tx":
            port_obj_list += ports_obj.P_PORT_PROTOTYPE
        elif direction_filter == "Rx":
            port_obj_list += ports_obj.R_PORT_PROTOTYPE
        else:
            port_obj_list += ports_obj.P_PORT_PROTOTYPE
            port_obj_list += ports_obj.R_PORT_PROTOTYPE

        for port_obj in port_obj_list:
            port_tag = port_obj.get_tag()
            if port_tag not in ["P-PORT-PROTOTYPE", "R-PORT-PROTOTYPE"]:
                continue
            if port_tag == "P-PORT-PROTOTYPE":
                nonqueued_com_spec_s = getattr(port_obj.PROVIDED_COM_SPECS, "NONQUEUED_SENDER_COM_SPEC", [])
                direction = "Tx"
                ports_key = "pports"
            elif port_tag == "R-PORT-PROTOTYPE":
                nonqueued_com_spec_s = getattr(port_obj.REQUIRED_COM_SPECS, "NONQUEUED_RECEIVER_COM_SPEC", [])
                direction = "Rx"
                ports_key = "rports"
            for nonqueued_com_spec in nonqueued_com_spec_s:
                port = {
                    "Tag": port_tag,
                    "DATA-ELEMENT-REF": {},
                    "Name": get_short_name(port_obj),
                    "Port_path": AutosarUtil.get_path(port_obj),
                    "Component": swc_name,
                    "Component_path": component_path,
                    "direction": direction,
                    "data_mappings": [],
                    "allowable_signal_type": "system_signal",
                }
                data_element_ref = getattr(getattr(nonqueued_com_spec, "DATA_ELEMENT_REF", None), "valueOf_", None)
                if data_element_ref is None:
                    continue
                idt_name = cls.__get_data_element_idt(data_element_ref)
                if idt_name not in idt_dict.keys():
                    continue
                if idt_dict[idt_name]["native_category"] == "STRUCTURE":
                    port["allowable_signal_type"] = "system_signal_group"
                port["DATA-ELEMENT-REF"] = {
                    "name": data_element_ref.split("/")[-1],
                    "ref": data_element_ref,
                    "idt": idt_name,
                    "idt_detail": idt_dict[idt_name],
                }
                swc_ports[swc_name][ports_key].append(port)
                all_port_data_s[f"{swc_name}/{port['Name']}/{data_element_ref.split('/')[-1]}"] = port
        swc_ports[swc_name]["pports"].sort(key=lambda x: x["Name"])
        swc_ports[swc_name]["rports"].sort(key=lambda x: x["Name"])
        return swc_ports, all_port_data_s

    @classmethod
    def _get_idt_original_type(cls, idt_dict, find_idt_name):
        rv = {"native_category_refs": [], "type_refs": []}
        if find_idt_name in cls.NATIVE_DATATYPES_SIZE.keys():
            rv["native_category_refs"].append("NATIVE_DATATYPE")
            rv["type_refs"].append({"native_category": "NUMERICAL", "basetype": find_idt_name})
            return rv
        basetype = find_idt_name
        counter = 0
        while basetype in idt_dict.keys():
            counter += 1
            if counter > 100:
                logger.error(f"Find idt '{find_idt_name}' loops exceeds the limit (100)!")
                break
            basetype_last = basetype
            basetype = idt_dict[basetype_last].get("basetype", None)
            native_category = idt_dict[basetype_last].get("native_category", None)
            rv["native_category_refs"].append(native_category)
            rv["type_refs"].append(idt_dict[basetype_last])
            if basetype in cls.NATIVE_DATATYPES_SIZE.keys():
                break
            if basetype == basetype_last:
                break
            if native_category not in ["NUMERICAL", "STRUCTURE", "ARRAY"]:
                break
        return rv

    @classmethod
    def update_idts_basetype(cls, idt_dict):
        # [ "NUMERICAL", "UNION", "STRUCTURE", "ARRAY"]
        idt_dict = deepcopy(idt_dict)
        for idt_name, idt in idt_dict.items():
            support_mapping = True
            native_category = idt.get("native_category", None)
            if native_category == "STRUCTURE":
                for member in idt.get("members", []):
                    find_info = cls._get_idt_original_type(idt_dict, member.get("datatype", None))
                    native_category_refs = find_info["native_category_refs"]
                    type_refs = find_info["type_refs"]
                    if (
                        "UNION" in native_category_refs
                        or "STRUCTURE" in native_category_refs
                        or native_category_refs.count("ARRAY") > 1
                        or len(type_refs) == 0
                    ):  # noqa:E129
                        support_mapping = False
                        break
                    if "name" in type_refs[-1].keys():
                        type_refs[-1].pop("name")
                    member.update(type_refs[-1])
            elif native_category in ["NUMERICAL", "ARRAY"]:
                find_info = cls._get_idt_original_type(idt_dict, idt.get("basetype", None))
                native_category_refs = find_info["native_category_refs"]
                type_refs = find_info["type_refs"]
                if (
                    "UNION" in native_category_refs
                    or "STRUCTURE" in native_category_refs
                    or "ARRAY" in native_category_refs
                    or len(type_refs) == 0
                ):  # noqa:E129
                    support_mapping = False
                else:
                    idt["basetype"] = type_refs[-1].get("basetype", None)
            else:
                support_mapping = False
            idt["support_data_mapping"] = support_mapping
        return idt_dict

    @classmethod
    def get_port_to_signal_pairs_from_flatextract(cls, flatextract):
        datamppings_info = {"sr_signal_mappings": {}, "sr_signal_group_mappings": {}}
        if flatextract is None:
            return datamppings_info
        flatextract_datamppings = SystemArxmlsManager.get_all_data_mappings(flatextract)
        for mapping_pair in flatextract_datamppings["sr_signal_mappings"]:
            component = mapping_pair["data_element_iref"]["context_component_ref"].split("/")[-1]
            port = mapping_pair["data_element_iref"]["context_port_ref"].split("/")[-1]
            data_prototype = mapping_pair["data_element_iref"]["target_data_prototype_ref"].split("/")[-1]
            swc_p_d = f"{component}/{port}/{data_prototype}"
            if swc_p_d not in datamppings_info["sr_signal_mappings"].keys():
                datamppings_info["sr_signal_mappings"][swc_p_d] = {
                    "data_element_iref": mapping_pair["data_element_iref"],
                    "system_singal_refs": [],
                    "type_mappings": {},
                }
            datamppings_info["sr_signal_mappings"][swc_p_d]["system_singal_refs"].append(
                mapping_pair["system_signal_ref"]
            )
        for mapping_pair in flatextract_datamppings["sr_signal_group_mappings"]:
            component = mapping_pair["data_element_iref"]["context_component_ref"].split("/")[-1]
            port = mapping_pair["data_element_iref"]["context_port_ref"].split("/")[-1]
            data_prototype = mapping_pair["data_element_iref"]["target_data_prototype_ref"].split("/")[-1]
            swc_p_d = f"{component}/{port}/{data_prototype}"
            if swc_p_d not in datamppings_info["sr_signal_group_mappings"].keys():
                datamppings_info["sr_signal_group_mappings"][swc_p_d] = {
                    "data_element_iref": mapping_pair["data_element_iref"],
                    "system_singal_refs": [],
                    "type_mappings": {},
                }
            datamppings_info["sr_signal_group_mappings"][swc_p_d]["system_singal_refs"].append(
                mapping_pair["signal_group_ref"]
            )
            datamppings_info["sr_signal_group_mappings"][swc_p_d]["type_mappings"].update(
                {mapping_pair["signal_group_ref"]: mapping_pair["type_mapping"]}
            )
        return datamppings_info

    @classmethod
    def get_all_signal_phys_props(cls, communication_autosar):
        rv = {}
        system_signals = communication_autosar.find_nodes("SYSTEM-SIGNAL")
        for system_signal in system_signals:
            short_name = get_short_name(system_signal)
            # Missing PHYSICAL-PROPS used to AttributeError and abort all verify.
            # Skip so validation can finish; callers treat absent keys as unconfigured.
            if system_signal.PHYSICAL_PROPS is None:
                continue
            variants = system_signal.PHYSICAL_PROPS.SW_DATA_DEF_PROPS_VARIANTS
            if variants is None or not variants.SW_DATA_DEF_PROPS_CONDITIONAL:
                continue
            conditional = variants.SW_DATA_DEF_PROPS_CONDITIONAL[0]
            compu_method_ref = get_valueOf(conditional.COMPU_METHOD_REF)
            data_constr_ref = get_valueOf(conditional.DATA_CONSTR_REF)
            rv.update({short_name: {"compu_method_ref": compu_method_ref, "data_constr_ref": data_constr_ref}})
        return rv

    @classmethod
    def update_cm_st_to_system_signals(cls, system_signals_dict, all_signal_phys_props, all_compu_methods):
        # update compu-method and signal base type
        for sig_name, _ in system_signals_dict["Tx"]["system_signal"].items():
            cls.get_cm_to_system_signal(system_signals_dict, all_signal_phys_props, sig_name, "Tx", all_compu_methods)
            cls.get_st_to_system_signal(system_signals_dict, sig_name, "Tx")

        for sig_name, _ in system_signals_dict["Rx"]["system_signal"].items():
            cls.get_cm_to_system_signal(system_signals_dict, all_signal_phys_props, sig_name, "Rx", all_compu_methods)
            cls.get_st_to_system_signal(system_signals_dict, sig_name, "Rx")

        for sig_group_name, _ in system_signals_dict["Tx"]["system_signal_group"].items():
            cls.get_cm_to_system_signal_group(
                system_signals_dict, all_signal_phys_props, sig_group_name, "Tx", all_compu_methods
            )
            cls.get_st_to_system_signal_group(system_signals_dict, sig_group_name, "Tx")

        for sig_group_name, _ in system_signals_dict["Rx"]["system_signal_group"].items():
            cls.get_cm_to_system_signal_group(
                system_signals_dict, all_signal_phys_props, sig_group_name, "Rx", all_compu_methods
            )
            cls.get_st_to_system_signal_group(system_signals_dict, sig_group_name, "Rx")

    @classmethod
    def get_cm_to_system_signal(cls, system_signals_dict, all_signal_phys_props, sig_name, dir, all_compu_methods):
        # 如果signal中包含了网段名称，那么先对完整signal的路径进行查找Compumethod
        # 再对去除网段信息的signal的路径进行查找Compumethod（适配达芬奇）
        act_sig_name = sig_name.split("_vcos_")[0] if "_vcos_" in sig_name else sig_name
        phys_props = all_signal_phys_props.get(sig_name, {})
        compu_method = {}
        system_signals_dict[dir]["system_signal"][sig_name].update({"compu_method": compu_method})
        if phys_props:
            compu_method_ref = phys_props.get("compu_method_ref", None)
            compu_method = all_compu_methods.get(compu_method_ref, {})
            if compu_method:
                system_signals_dict[dir]["system_signal"][sig_name].update({"compu_method": compu_method})
            else:
                if act_sig_name != sig_name:
                    temp = compu_method_ref.split("/")[0:-1]
                    temp.append(f"CM_{act_sig_name}")
                    compu_method_ref = "/".join(temp)
                    compu_method = all_compu_methods.get(compu_method_ref, {})
                    system_signals_dict[dir]["system_signal"][sig_name].update({"compu_method": compu_method})

    @classmethod
    def get_cm_to_system_signal_group(
        cls, system_signals_dict, all_signal_phys_props, sig_group_name, dir, all_compu_methods
    ):
        compu_method = {}
        system_signals = system_signals_dict[dir]["system_signal_group"][sig_group_name]["system_signal"]
        for sig_name in system_signals.keys():
            act_sig_name = sig_name.split("_vcos_")[0] if "_vcos_" in sig_name else sig_name
            phys_props = all_signal_phys_props.get(sig_name, {})
            system_signals[sig_name]["compu_method"] = {}
            if not phys_props:
                continue
            compu_method_ref = phys_props.get("compu_method_ref", None)
            compu_method = all_compu_methods.get(compu_method_ref, {})
            if compu_method:
                system_signals[sig_name]["compu_method"] = compu_method
            else:
                if act_sig_name != sig_name:
                    temp = compu_method_ref.split("/")[0:-1]
                    temp.append(f"CM_{act_sig_name}")
                    compu_method_ref = "/".join(temp)
                    compu_method = all_compu_methods.get(compu_method_ref, {})
                    system_signals[sig_name]["compu_method"] = compu_method

    @classmethod
    def get_st_to_system_signal(cls, system_signals_dict, sig_name, dir):
        sig_type = system_signals_dict[dir]["system_signal"][sig_name]["detail_info"]["ComSignalType"]
        system_signals_dict[dir]["system_signal"][sig_name].update({"real_sig_type": sig_type.lower()})

    @classmethod
    def get_st_to_system_signal_group(cls, system_signals_dict, sig_group_name, dir):
        system_signals = system_signals_dict[dir]["system_signal_group"][sig_group_name]["system_signal"]
        for sig_name, sig_value in system_signals.items():
            sig_type = sig_value["detail_info"]["ComSignalType"]
            system_signals[sig_name].update({"real_sig_type": sig_type.lower()})

    @classmethod
    def get_system_signal_to_isignal_mappings_info(cls, communication_autosar, com_autosar, all_compu_methods={}):
        system_signals_dict = SystemArxmlsManager.get_communication_all_system_signal_mapping_info(
            communication_autosar
        )
        com_signals_info = cls.get_com_signals_info(com_autosar)
        all_signal_phys_props = cls.get_all_signal_phys_props(communication_autosar)
        system_signals_dict = cls.update_com_signal_info_to_system_signals(system_signals_dict, com_signals_info)
        if all_compu_methods:
            cls.update_cm_st_to_system_signals(system_signals_dict, all_signal_phys_props, all_compu_methods)
        return system_signals_dict

    @classmethod
    def get_com_signals_info(cls, com_autosar):

        def _parse_container_to_dict(obj):
            short_name = get_short_name(obj)
            container_dict = {"short_name": short_name}
            para_obj_list = getattr(obj.PARAMETER_VALUES, "ECUC_NUMERICAL_PARAM_VALUE", []) + getattr(
                obj.PARAMETER_VALUES, "ECUC_TEXTUAL_PARAM_VALUE", []
            )
            for para_obj in para_obj_list:
                definition = get_definition_ref(para_obj).split("/")[-1]
                container_dict.update({definition: get_value(para_obj)})
            for para_obj in getattr(obj.REFERENCE_VALUES, "ECUC_REFERENCE_VALUE", []):
                definition = get_definition_ref(para_obj).split("/")[-1]
                if definition not in container_dict.keys():
                    container_dict[definition] = []
                container_dict[definition].append(get_value_ref(para_obj))
            return container_dict

        com_signals_info = {}
        if com_autosar is None:
            return {}
        com_config_obj = com_autosar.find("/ActiveEcuC/Com/ComConfig")
        ecuc_containers = getattr(getattr(com_config_obj, "SUB_CONTAINERS", None), "ECUC_CONTAINER_VALUE", [])
        for container in ecuc_containers:
            if not str(get_definition_ref(container)).endswith("ComIPdu"):
                continue
            com_ipdu_dict = _parse_container_to_dict(container)
            com_tx_mode = container.find("./ComTxIPdu/ComTxModeTrue/ComTxMode")
            if com_tx_mode is not None:
                com_tx_mode_dict = _parse_container_to_dict(com_tx_mode)
            signal_ref_s = com_ipdu_dict.get("ComIPduSignalRef", []) + com_ipdu_dict.get("ComIPduSignalGroupRef", [])
            if not isinstance(signal_ref_s, list):
                continue
            for signal_ref in signal_ref_s:
                if signal_ref in [None, ""]:
                    continue
                signal_container = com_autosar.find(signal_ref)
                if signal_container is None:
                    logger.warning(f'"{signal_ref}" not found in COM!')
                    continue
                signal_info_dict = {"detail_info": _parse_container_to_dict(signal_container)}
                if com_tx_mode is not None:
                    signal_info_dict["com_tx_mode"] = com_tx_mode_dict
                if get_definition_ref(signal_container).split("/")[-1] == "ComSignalGroup":
                    signal_info_dict["com_group_signal"] = {}
                    groupsignal_containers = getattr(
                        getattr(signal_container, "SUB_CONTAINERS", None), "ECUC_CONTAINER_VALUE", []
                    )
                    for groupsignal in groupsignal_containers:
                        shortname = get_short_name(groupsignal)
                        signal_info_dict["com_group_signal"][shortname] = _parse_container_to_dict(groupsignal)
                shortname = get_short_name(signal_container)
                com_signals_info[shortname] = signal_info_dict
        return com_signals_info

    @classmethod
    def update_com_signal_info_to_system_signals(cls, system_signals_dict, com_signals_info):
        for _, signals_signalgroups_one_direction in system_signals_dict.items():
            for signal_type, signals_signalgroups in signals_signalgroups_one_direction.items():
                for system_signal_name in list(signals_signalgroups.keys()):
                    signal_signalgroup = signals_signalgroups[system_signal_name]
                    signal_siggroup_name_guess = f'{signal_signalgroup["full_name"]}_{signal_signalgroup["direction"]}'
                    if signal_siggroup_name_guess not in com_signals_info.keys():
                        logger.warning(f"Signal {signal_siggroup_name_guess} not found in COM!")
                        signals_signalgroups.pop(system_signal_name)
                        continue
                    if signal_type == "system_signal_group":
                        for _, groupsignal in signal_signalgroup["system_signal"].items():
                            groupsignal_name_guess = f'{groupsignal["full_name"]}_{groupsignal["direction"]}'
                            if (
                                groupsignal_name_guess
                                not in com_signals_info[signal_siggroup_name_guess]["com_group_signal"].keys()
                            ):
                                signals_signalgroups.pop(system_signal_name)
                                logger.warning(f"GroupSignal {groupsignal_name_guess} not found in COM!")
                                break
                            groupsignal["detail_info"] = com_signals_info[signal_siggroup_name_guess][
                                "com_group_signal"
                            ][groupsignal_name_guess]
                        else:
                            signal_signalgroup.update(com_signals_info[signal_siggroup_name_guess])
                    else:
                        signal_signalgroup.update(com_signals_info[signal_siggroup_name_guess])
        return system_signals_dict

    @classmethod
    def update_data_mappings_to_ports(cls, ports_dict, data_mappings_dict, system_signal_isignals):
        for swc_p_d, system_singal_refs in data_mappings_dict["sr_signal_mappings"].items():
            port_info = ports_dict.get(swc_p_d, None)
            if port_info is None or port_info["direction"] not in ["Tx", "Rx"]:
                continue
            port_info["data_mappings"] = []
            sig_type = port_info["allowable_signal_type"]
            for system_singal_ref in system_singal_refs["system_singal_refs"]:
                system_singal_name = system_singal_ref.split("/")[-1]
                signal_info = deepcopy(
                    system_signal_isignals[port_info["direction"]][sig_type].get(system_singal_name, None)
                )
                if signal_info is None:
                    continue
                signal_info["type_compat"] = cls.get_signal_idt_type_compatibility(
                    {"signal": signal_info}, port_info["DATA-ELEMENT-REF"]["idt_detail"]
                )
                port_info["data_mappings"].append(signal_info)
        for swc_p_d, system_singal_refs in data_mappings_dict["sr_signal_group_mappings"].items():
            port_info = ports_dict.get(swc_p_d, None)
            if port_info is None or port_info["direction"] not in ["Tx", "Rx"]:
                continue
            port_info["data_mappings"] = []
            port_info["type_mappings"] = {}
            sig_type = port_info["allowable_signal_type"]
            for system_singal_ref in system_singal_refs["system_singal_refs"]:
                system_singal_name = system_singal_ref.split("/")[-1]
                signal_info = deepcopy(
                    system_signal_isignals[port_info["direction"]][sig_type].get(system_singal_name, None)
                )
                if signal_info is None:
                    continue
                idt_detail = port_info["DATA-ELEMENT-REF"]["idt_detail"]
                signal_info["type_compat"] = cls.get_signal_idt_type_compatibility({"signal": signal_info}, idt_detail)
                full_name = signal_info["full_name"]
                port_info["data_mappings"].append(signal_info)
                port_info["type_mappings"].update({full_name: system_singal_refs["type_mappings"][system_singal_ref]})
                for type_mapping in port_info["type_mappings"][full_name]:
                    system_signal = type_mapping["system_signal"]
                    i_signal_name = signal_info["system_signal_to_i_signal"].get(system_signal)
                    type_mapping["i_signal"] = i_signal_name
                    system_signal_info = signal_info["system_signal"][system_signal]
                    record_element_info = cls.__get_record_element(
                        idt_detail, type_mapping["implementation_record_element"]
                    )
                    if record_element_info is None:
                        continue
                    type_mapping["type_compat"] = cls.get_signal_idt_type_compatibility(
                        {"signal": system_signal_info}, record_element_info
                    )
        return ports_dict

    @classmethod
    def __get_record_element(cls, idt, element_name):
        for member_info in idt.get("members", []):
            if member_info["name"] == element_name:
                return member_info
        return None

    @classmethod
    def ports_to_signal_mapping_recommend(cls, ports: list, system_signals: dict, conditions: dict):
        system_signals_pre_processing = cls._system_signals_pre_processing(system_signals, conditions)
        for port in ports:
            need_mapping = cls._port_filter_result(port, conditions)
            if need_mapping is True:
                port["signals_recommend"] = cls._port_to_signal_mapping_recommend(
                    port, system_signals_pre_processing, conditions
                )
            else:
                port["signals_recommend"] = []
        return ports

    @classmethod
    def _port_filter_result(cls, port, filter_condition):
        if filter_condition.get("case_sensitive", False) is True:
            name_temp = port["Name"]
        else:
            name_temp = port["Name"].lower()
        rv = {"need_mapping": False}
        if filter_condition is None or filter_condition == {}:
            rv["need_mapping"] = True
            port["port_name_to_mapping"] = name_temp
        elif (
            filter_condition["direction"] == port["direction"]
            and port["Name"].startswith(filter_condition["prefix"])
            and port["Name"].endswith(filter_condition["postfix"])
        ):
            rv["need_mapping"] = True
            start_idx = len(filter_condition["prefix"])
            end_idx = len(name_temp) - len(filter_condition["postfix"])
            port["port_name_to_mapping"] = name_temp[start_idx:end_idx]
        else:
            rv["need_mapping"] = False
            port["port_name_to_mapping"] = name_temp
        return rv["need_mapping"]

    @classmethod
    def _port_to_signal_mapping_recommend(cls, port_info: dict, system_signals: dict, conditions: dict):
        if conditions.get("case_sensitive") is True:
            port_name = port_info["port_name_to_mapping"]
            data_name = port_info["DATA-ELEMENT-REF"]["name"]
        else:
            port_name = port_info["port_name_to_mapping"].lower()
            data_name = port_info["DATA-ELEMENT-REF"]["name"].lower()
        direction = port_info["direction"]
        signal_type = port_info["allowable_signal_type"]
        signals_recommend = cls.check_data_mapping(
            port_name, data_name, system_signals[direction][signal_type], conditions
        )
        idt_detail = port_info.get("DATA-ELEMENT-REF", {}).get("idt_detail", None)
        if idt_detail is None:
            return []
        cls._update_compatibility_data_type(signals_recommend, idt_detail)
        if signal_type == "system_signal_group" and len(signals_recommend) > 0:
            signalgroups_recommend = []
            for signal_group in signals_recommend:
                signal_group["signal"]["type_mappings_brief"] = {}
                data_members = port_info["DATA-ELEMENT-REF"]["idt_detail"]["members"]
                for data_member in data_members:
                    group_signal = cls.check_data_mapping(
                        "/", data_member["name"], signal_group["signal"]["system_signal"], conditions
                    )
                    if len(group_signal) < 1 or group_signal is None:
                        break
                    # idt_info = cls._record_element_data_type_parse(data_member)
                    cls._update_compatibility_data_type(group_signal, data_member)
                    if signal_group["compatibility"].value[0] < group_signal[0]["compatibility"].value[0]:
                        signal_group["compatibility"] = group_signal[0]["compatibility"]
                    signal_group["signal"]["type_mappings_brief"].update(
                        {
                            data_member["name"]: cls._get_type_mappings_brief(
                                port_info["Name"], signal_group, group_signal[0], data_member["name"]
                            )
                        }
                    )
                members_num = len(port_info["DATA-ELEMENT-REF"]["idt_detail"]["members"])
                group_signal_num = len(signal_group["signal"]["system_signal"])
                if group_signal_num == members_num:
                    signalgroups_recommend.append(signal_group)
            signals_recommend = signalgroups_recommend

        return signals_recommend

    @classmethod
    def _record_element_data_type_parse(cls, data_member):
        idt_info = None
        if data_member.get("category", None) == "VALUE":
            idt_info = {"native_category": "NUMERICAL", "basetype": data_member.get("datatype", None)}
        elif data_member.get("category", None) == "ARRAY":
            idt_dict_info = cls.idt_dict.get(data_member.get("datatype", None), {})
            idt_info = {
                "native_category": "ARRAY",
                "basetype": idt_dict_info.get("basetype", None),
                "size": idt_dict_info.get("size", 0),
            }
        elif data_member.get("category", None) == "TYPE_REFERENCE":
            idt_dict_info = cls.idt_dict.get(data_member.get("datatype", None), {})
            if idt_dict_info.get("native_category", None) != "NUMERICAL":
                return idt_info
            idt_info = {"native_category": "NUMERICAL", "basetype": data_member.get("datatype", None)}
        elif data_member.get("category", None) == "STRUCTURE":
            pass
        return idt_info

    @classmethod
    def _update_compatibility_data_type(cls, signals_recommend, port_idt_detail, is_record_element=False):
        for signal in signals_recommend:
            data_type_check_result = cls.get_signal_idt_type_compatibility(signal, port_idt_detail)
            if data_type_check_result == Compatibility.full:  # 类型一致不更新兼容性
                pass
            elif data_type_check_result == Compatibility.partial:
                if signal["compatibility"] == Compatibility.full:
                    signal["compatibility"] = Compatibility.full_datalength_compatible
                elif signal["compatibility"] == Compatibility.partial:
                    signal["compatibility"] = Compatibility.partial_datalength_compatible
                else:
                    logger.warning("Data Mapping Compatibility not found!")
            elif data_type_check_result == Compatibility.datalength_imcompatible:
                signal["compatibility"] = Compatibility.datalength_imcompatible
            else:
                logger.warning("Unexpected branch!")

    @classmethod
    def get_signal_idt_type_compatibility(cls, signal, port_idt_detail):
        native_category = port_idt_detail.get("native_category", None)
        data_type_check_result = Compatibility.datalength_imcompatible
        signal_category = signal["signal"]["signal_type"]
        com_signal_type = signal["signal"]["detail_info"].get("ComSignalType", None)
        com_signal_length_str = signal["signal"]["detail_info"].get("ComSignalLength", None)
        port_data_match_list = cls.DATA_TYPE_MATCHING.get(com_signal_type, [])
        if native_category == "NUMERICAL" and com_signal_type is not None and str(com_signal_length_str).isdigit():
            com_signal_length = int(com_signal_length_str)
            port_data_basetype = port_idt_detail["basetype"]
            port_data_basetype_size = cls.NATIVE_DATATYPES_SIZE.get(port_data_basetype, 0)
            if port_data_basetype_size == 0:
                logger.warning(
                    f'Data"{port_idt_detail.get("name", "")}" not support AutoDataMapping, '
                    f'Because type"{port_data_basetype}" not found!'
                )
            if port_data_basetype in port_data_match_list:  # 类型一致
                data_type_check_result = Compatibility.full
            elif port_data_basetype_size >= com_signal_length:  # 类型兼容
                data_type_check_result = Compatibility.partial
            else:  # 不兼容
                data_type_check_result = Compatibility.datalength_imcompatible
        elif native_category == "ARRAY" and com_signal_type is not None and str(com_signal_length_str).isdigit():
            com_signal_length = int(com_signal_length_str)
            port_data_basetype = port_idt_detail["basetype"]
            port_data_size = port_idt_detail["size"]
            port_data_basetype_size = cls.NATIVE_DATATYPES_SIZE.get(port_data_basetype, 0)
            if port_data_basetype_size == 0:
                logger.warning(
                    f'Data"{port_idt_detail.get("name", "")}" not support AutoDataMapping, '
                    f'Because type"{port_data_basetype}" not found!'
                )
            port_data_all_byte_size = port_data_basetype_size * port_data_size
            if port_data_basetype in cls.DATA_TYPE_MATCHING["UINT8"] and com_signal_type == "UINT8_N":
                if port_data_size == com_signal_length:  # 类型一致
                    data_type_check_result = Compatibility.full
                elif port_data_size > com_signal_length:  # 类型兼容
                    data_type_check_result = Compatibility.partial
                else:  # 不兼容
                    data_type_check_result = Compatibility.datalength_imcompatible
            elif port_data_all_byte_size >= com_signal_length:  # 类型长度兼容
                data_type_check_result = Compatibility.partial
            else:  # 类型长度不兼容
                data_type_check_result = Compatibility.datalength_imcompatible
        elif native_category == "STRUCTURE" and signal_category == "system_signal_group":
            data_type_check_result = Compatibility.full
        return data_type_check_result

    @classmethod
    def _system_signals_pre_processing(cls, system_signals: dict, conditions: dict):
        directions = ["Tx", "Rx"]
        signal_types = ["system_signal", "system_signal_group"]
        rv = {}
        if conditions.get("case_sensitive") is True:
            return system_signals
        else:
            for direction in directions:
                rv[direction] = {}
                for sig_type in signal_types:
                    rv[direction][sig_type] = {
                        key.lower(): signal for key, signal in system_signals[direction][sig_type].items()
                    }
        return rv

    @classmethod
    def system_signals_filter(cls, system_signals, conditions):
        direction = conditions["direction"]
        signal_type = conditions["signal_type"]
        return system_signals.get(direction, {}).get(signal_type, [])

    @classmethod
    def excel_to_dict_list(cls, file_path):
        # 将Excel文件一个工作表转换为字典列表，取出第一行(去除前后的空格)作为字典的key，其余行(去除空格)作为字典的value
        data = []
        if not os.path.isfile(file_path):
            raise FileNotFoundError(f"The file {file_path} does not exist.")
        if not file_path.endswith((".xlsx", ".xlsm")):
            raise ValueError(f"The file {file_path} is not a valid Excel file.")
        try:
            workbook = load_workbook(filename=file_path)
        except Exception as e:
            raise ValueError(f"Failed to read Excel file {file_path}: {e}")
        sheet = workbook[workbook.sheetnames[0]]
        keys = [
            cell.value.strip() if isinstance(cell.value, str) else cell.value
            for cell in next(sheet.iter_rows(min_row=1, max_row=1))
        ]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            row_values = [value.replace(" ", "") if isinstance(value, str) else value for value in row]
            record = dict(zip(keys, row_values))
            data.append(record)
        return data

    @classmethod
    def add_data_mapping_to_flatextract_by_file(cls, autosar_system_class, file_data_s):
        for data in file_data_s:
            auto_mappings = data.get("auto_mappings", None)
            if not auto_mappings:
                continue
            ret_list = cls.add_new_data_mappings_to_arxml(autosar_system_class, auto_mappings)
            if False in ret_list:
                data["mapping_result"] = {"result": False, "reason": "add new data mapping error"}
            else:
                data["mapping_result"] = {"result": True, "reason": ""}

    @classmethod
    def update_auto_mapping_dict_by_file(cls, file_data_s):
        all_mapping_dicts = []
        for data in file_data_s:
            if not data.get("mapping_result", {}).get("result") or "." in data[_element_key]:
                continue
            data["auto_mappings"] = []
            data["find_pdu"] = [signal["signal"]["pdu"] for signal in data["port"].get("signals_recommend", [])]
            mapping_dicts = DataMapping.get_sr_to_signal_mapping_dicts([data["port"]])
            for mapping_dict in mapping_dicts:
                if mapping_dict in all_mapping_dicts:
                    continue
                all_mapping_dicts.append(mapping_dict)
                data["auto_mappings"].append(mapping_dict)

    @classmethod
    def auto_mapping_file_data_preprocess(cls, row_data_s, datamapping_infos):
        group_mapping_data_s = []
        group_mapping = None
        all_swc_ports = datamapping_infos[cls.INFOS_INDEX.SWC_PORTS]
        system_signals_dict = datamapping_infos[cls.INFOS_INDEX.SYSTEM_SIGNALS]
        for data in row_data_s:
            if not data[_signal_key]:
                continue
            swc_p_d = f"{data[_component_key]}/{data[_port_key]}/{str(data[_element_key]).split('.')[0]}"
            data.update({"port": deepcopy(all_swc_ports["all_port_data_s"].get(swc_p_d, None))})
            if data["port"] is None:
                data["mapping_result"] = {"result": False, "reason": "port not found"}
                continue
            data["port"]["signals_recommend"] = []
            signal_type = data["port"]["allowable_signal_type"]
            if signal_type == "system_signal_group":
                if "." not in data[_element_key]:
                    data.update({"submembers": [], "mapping_result": {"result": True, "reason": ""}})
                    group_mapping = data
                    group_mapping_data_s.append(group_mapping)
                else:
                    if group_mapping and data["port"] == group_mapping["port"]:
                        data["mapping_result"] = group_mapping["mapping_result"]
                        group_mapping["submembers"].append(data)
                continue
            else:
                cls.update_signal_mapping_data_by_file(data, system_signals_dict)
        for group_mapping in group_mapping_data_s:
            if len(group_mapping["submembers"]) > 0:
                cls.update_group_mapping_data_by_file(group_mapping, system_signals_dict)
            else:
                group_mapping["mapping_result"] = {"result": False, "reason": "structure no member"}

    @classmethod
    def update_signal_mapping_data_by_file(cls, data, system_signals_dict):
        direction = data["port"]["direction"]
        signal_type = data["port"]["allowable_signal_type"]
        _signals_recommend = cls.check_data_mapping(
            data[_signal_key],
            data[_signal_key],
            system_signals_dict[direction][signal_type],
            {"allow_multi_signal": True, "whole_word": True},
        )
        _signals_recommend = [
            signal
            for signal in _signals_recommend
            if (
                (not data[_pdu_key] or signal["signal"]["pdu"] in data[_pdu_key].split(","))
                and signal["compatibility"] == Compatibility.full
            )
        ]
        if len(_signals_recommend) < 1:
            data["mapping_result"] = {"result": False, "reason": "signal not found"}
            return
        if data[_pdu_key] and len(_signals_recommend) < len(data[_pdu_key].split(",")):
            data["mapping_result"] = {"result": False, "reason": "pdu not found"}
            return
        data["mapping_result"] = {"result": True, "reason": ""}
        for signal in _signals_recommend:
            signal.update({"auto_mapping": True})
            for mapped in data["port"]["data_mappings"]:
                if mapped["short_name"] == signal["signal"]["short_name"]:
                    signal.update({"auto_mapping": False, "reason": "already mapped"})
                    break
        data["port"]["signals_recommend"] += _signals_recommend

    @classmethod
    def update_group_mapping_data_by_file(cls, group_mapping, system_signals_dict):
        type_member_signal_info = {}
        direction = group_mapping["port"]["direction"]
        signal_type = group_mapping["port"]["allowable_signal_type"]
        _signals_recommend = cls.check_data_mapping(
            group_mapping[_signal_key],
            group_mapping[_signal_key],
            system_signals_dict[direction][signal_type],
            {"allow_multi_signal": True, "whole_word": True},
        )
        _signals_recommend = [
            signal
            for signal in _signals_recommend
            if (
                (not group_mapping[_pdu_key] or signal["signal"]["pdu"] in group_mapping[_pdu_key].split(","))
                and signal["compatibility"] == Compatibility.full
            )
        ]
        if len(_signals_recommend) < 1:
            group_mapping["mapping_result"].update({"result": False, "reason": "signalgroup not found"})
        for data in group_mapping["submembers"]:
            type_member_signal_info.update({data[_element_key].split(".")[-1]: data[_signal_key].split(".")[-1]})
        for signal_group in _signals_recommend:
            signal_group["signal"]["type_mappings_brief"] = {}
            data_members = group_mapping["port"]["DATA-ELEMENT-REF"]["idt_detail"]["members"]
            data_member_names = [mem["name"] for mem in data_members]
            for member_name, signal_name in type_member_signal_info.items():
                if member_name not in data_member_names:
                    group_mapping["mapping_result"].update({"result": False, "reason": "member not found"})
                    return
                group_signal = cls.check_data_mapping(
                    "/",
                    signal_name,
                    signal_group["signal"]["system_signal"],
                    {"allow_multi_signal": True, "whole_word": True},
                )
                if len(group_signal) < 1:
                    group_mapping["mapping_result"].update({"result": False, "reason": "groupsignal not found"})
                    return
                signal_group["signal"]["type_mappings_brief"].update(
                    {
                        member_name: cls._get_type_mappings_brief(
                            group_mapping["port"]["Name"], signal_group, group_signal[0], member_name
                        )
                    }
                )
            signal_group.update({"auto_mapping": True})
            for mapped in group_mapping["port"]["data_mappings"]:
                if mapped["short_name"] == signal_group["signal"]["short_name"]:
                    signal_group.update({"auto_mapping": False, "reason": "already mapped"})
                    break
            group_mapping["port"]["signals_recommend"].append(signal_group)

    @classmethod
    def _get_type_mappings_brief(cls, port_name, signal_group, group_signal, member_name):
        return {
            "port_name": port_name,
            "implementation_record_element": member_name,
            "group_name": signal_group["signal"]["short_name"],
            "system_signal": group_signal["signal"]["short_name"],
        }

    @classmethod
    def auto_data_mapping_by_file(cls, autosar_asw_class, autosar_bsw_class, autosar_system_class, file_path):
        file_data_s = cls.excel_to_dict_list(file_path)
        if not file_data_s:
            logger.error(f"File: {file_path} no data!")
            return
        item_keys = list(file_data_s[0].keys())
        for item_name in [_element_key, _component_key, _port_key, _signal_key, _pdu_key]:
            if item_name not in item_keys:
                logger.error(f"File: {file_path} header line error: {item_name} not found!")
                return

        datamapping_infos = cls.get_datamapping_infos(autosar_asw_class, autosar_bsw_class, autosar_system_class, None)
        # 根据Excel数据和DataMapping信息，对Excel数据进行预处理，找到对应的Port和Signal添加到signals_recommend字段中
        cls.auto_mapping_file_data_preprocess(file_data_s, datamapping_infos)
        # 根据signals_recommend信息，更新Excel数据中的auto_mappings字段，为添加DataMapping连接做准备
        cls.update_auto_mapping_dict_by_file(file_data_s)
        # 根据signals_recommend信息，添加DataMapping连接关系到FlatExtract中
        cls.add_data_mapping_to_flatextract_by_file(autosar_system_class, file_data_s)
        return file_data_s

    @classmethod
    def data_mapping_info_to_codegen(cls, asw_autosar, system_autosar_dict, com_autosar):
        if not asw_autosar:
            return []

        flatextract_datamppings = DataMapping.get_port_to_signal_pairs_from_flatextract(
            system_autosar_dict.get("FlatExtract", None)
        )
        system_signals_dict = DataMapping.get_system_signal_to_isignal_mappings_info(
            system_autosar_dict.get("Communication", None), com_autosar, AutosarUtils.get_all_compu_methods(asw_autosar)
        )
        return [flatextract_datamppings, system_signals_dict]
